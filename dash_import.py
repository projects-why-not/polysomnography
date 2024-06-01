import random
from datetime import datetime
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from dash import Dash, dcc, html, Input, Output, callback
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx2pdf import convert
def set_cell_border(cell, **kwargs):
    """
    Set cell's border
    Usage:
    set_cell_border(
        cell,
        top={"sz": 12, "val": "single", "color": "000000", "space": "0"},
        bottom={"sz": 12, "val": "single", "color": "000000", "space": "0"},
        left={"sz": 12, "val": "single", "color": "000000", "space": "0"},
        right={"sz": 12, "val": "single", "color": "000000", "space": "0"},
    )
    """
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()

    for border_name in ["top", "left", "bottom", "right"]:
        border = kwargs.get(border_name)
        if border:
            element = OxmlElement(f"w:{border_name}")
            for attr, value in border.items():
                element.set(qn(f"w:{attr}"), str(value))
            tcPr.append(element)

def create_polysomnography_report(template_data, output_word_path, output_pdf_path):
    # Create a new Document
    doc = Document()

    # Title
    title = doc.add_heading('Отчет для врача ... по полисомнографической записи ...', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add technical report section
    doc.add_heading('Технический отчет', level=2)
    doc.add_paragraph(f'Проведен анализ ПСГ длительностью {template_data["duration"]} минут')
    doc.add_paragraph(f'Оценивались сигналы типа ЭЭГ')
    doc.add_paragraph(f'Время обработки составило {template_data["processing_time"]} секунд')

    # Add clinical report section
    doc.add_heading('Клинический отчет', level=2)
    doc.add_paragraph(f'По результатам анализа выявлено {template_data["breathing_disorder_presence"]} признаков нарушений дыхания во сне')
    doc.add_paragraph(f'За время ПСГ выявлено {template_data["breathing_disorder_episodes_count"]} эпизодов нарушения дыхания (НРД)')

    # Add table
    table = doc.add_table(rows=1, cols=5)
    hdr_cells = table.rows[0].cells
    headers = ['№ эпизода НРД', 'Время начала регистрации эпизода НРД, с', 'Время завершения регистрации НРД, с', 'Длительность эпизода НРД, с', 'Тип эпизода НРД']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_border(hdr_cells[i], top={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                      bottom={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                      left={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                      right={"sz": 12, "val": "single", "color": "000000", "space": "0"})

    for episode in template_data["episodes"]:
        row_cells = table.add_row().cells
        row_cells[0].text = str(episode["number"])
        row_cells[1].text = str(episode["start_time"])
        row_cells[2].text = str(episode["end_time"])
        row_cells[3].text = str(episode["duration"])
        row_cells[4].text = episode["type"]
        for cell in row_cells:
            set_cell_border(cell, top={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                  bottom={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                  left={"sz": 12, "val": "single", "color": "000000", "space": "0"},
                                  right={"sz": 12, "val": "single", "color": "000000", "space": "0"})

    # Save the document
    doc.save(output_word_path)

    # Convert the document to PDF
    convert(output_word_path, output_pdf_path)

def div_maker(fig):
    return html.Div(
        children=dcc.Graph(figure=fig),
        style={"display": "flex", "justify-content": "center"},
    )

def upper_plot_generator(data_poly, data_gipno):

    i = 0

    fig = go.Figure()
    for column in data_poly.columns[:-1]:
        if column in ['Airflow', 'Chest', 'Abdomen']:
            fig.add_trace(go.Scatter(x=data_poly['Time'], y=data_poly[column] + 0.005 * i, name=column))
            i += 1

    if data_gipno is not None:
        fig.add_trace(go.Scatter(x=data_gipno['Time'], y=data_gipno['stage'] / 900 + 0.005 * i, name='gipno'))

    fig.update_yaxes(showticklabels=False)
    fig.update_xaxes(
        tickformat="%H:%M:%S"
    )
    fig.update_xaxes(rangeslider_visible=True)  # ranges slider

    return div_maker(fig)

def lower_plot_generator(data):
    i = 0

    fig = go.Figure()
    for column in data.columns[:-1]:
        if column in ['Fp1-M2', 'C3-M2', 'O1-M2', 'Fp2-M1', 'C4-M1', 'O2-M1']:
            fig.add_trace(go.Scatter(x=data['Time'], y=data[column] + 0.0005 * i, name=column))
            i += 1
    fig.update_yaxes(showticklabels=False)
    fig.update_xaxes(
        tickformat="%H:%M:%S"
    )
    fig.update_xaxes(rangeslider_visible=True)  # ranges slider

    return div_maker(fig)

def start_dash(data_poly, data_gipno, hertz, processing_time, breathing_disorder_presence, path_word, path_pdf):
    data_poly, data_gipno, duration = data_preprocessing(data_poly, data_gipno, hertz)

    app = Dash(__name__)
    app._favicon = "images/icon3.png"
    app.title = "Результат анализа"

    title_page = html.H1(children='Результаты анализа сна по ЭЭГ', className='title')

    link_1 = html.A(children='Загрузить отчёт Word', href='link/to/your/download/file', download=path_word)
    link_2 = html.A(children='Загрузить отчёт PDF', href='link/to/your/download/file', download=path_pdf)

    # preloader = html.Div(className='preloader')

    # external_stylesheets = [
    #     {
    #         "rel": "stylesheet",
    #     },
    # ]

    app.layout = html.Div(children=[
        title_page,
        upper_plot_generator(data_poly, data_gipno),
        lower_plot_generator(data_poly),
        link_1,
        link_2],
        className='page')

    app.run_server(debug=True)

def data_preprocessing(data_poly, data_gipno, hertz):
    poly = pd.DataFrame()
    for column in data_poly.columns:
        poly[column] = data_poly[column].rolling(window=hertz).mean()[::hertz]
    poly.loc[0, poly.columns] = data_poly.loc[0, poly.columns]
    poly['Time'] = pd.date_range(start='00:00:00', periods=len(poly), freq='1s')


    data_gipno['Time'] = pd.date_range(start='00:00:00', periods=len(data_gipno), freq='30s')
    data_gipno = data_gipno[['Time', 'stage']].copy()

    duration = int((poly['Time'].iloc[-1] - poly['Time'].iloc[0]).total_seconds())

    return poly, data_gipno, duration

if __name__ == "__main__":
    data1 = pd.read_csv('full_data.csv')
    data2 = pd.read_csv('hypno.csv')
    # input
    hertz = 200

    processing_time = random.randrange(5,15)
    breathing_disorder_presence = 'наличие'
    # episodes_data = pd.read_csv('episodes_data.csv')
    start_dash(data1, data2, hertz, processing_time, breathing_disorder_presence, 'output_report.docx','output_report.pdf')


